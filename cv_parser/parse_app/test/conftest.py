import os
import pytest
import json
from docx import Document
from fpdf import FPDF
from tempfile import NamedTemporaryFile
from pytest_mock import mocker
from io import BytesIO
from unittest.mock import MagicMock
from pytest_mock import MockFixture


FILE_CONTENT : str = "Sample Document"

@pytest.fixture
def generate_pdf_file():
    with NamedTemporaryFile(delete=False, suffix=".pdf", mode="wb") as temp_file:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        pdf.cell(200, 10, txt=FILE_CONTENT, align="C")
        pdf.output(temp_file.name)
        print("file name", temp_file.name)
    yield temp_file.name

    os.remove(temp_file.name)

@pytest.fixture
def generate_doc_file():
    with NamedTemporaryFile(delete=False, suffix=".docx", mode="wb") as temp_file:
        doc = Document()
        doc.add_paragraph(FILE_CONTENT)
        doc.save(temp_file.name)
        
    yield temp_file.name

    os.remove(temp_file.name)

@pytest.fixture
def mock_invoke_model(mocker: MockFixture) -> MagicMock:
    sample_response = {
        "generation": '{"applicant_name": "Gino"}'
    }

    mock_response_get = BytesIO(json.dumps(sample_response).encode("utf-8"))

    invoke_model_response = mocker.MagicMock()
    invoke_model_response.get.return_value = mock_response_get
    invoke_model_mock_version = mocker.MagicMock()
    boto3_response = mocker.MagicMock()

    boto3_response.invoke_model = invoke_model_mock_version
    boto3_response.invoke_model.return_value = invoke_model_response

    mock_boto3_client = mocker.patch("boto3.client")
    mock_boto3_client.return_value = boto3_response

    return invoke_model_mock_version